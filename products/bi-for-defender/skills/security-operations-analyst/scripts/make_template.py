"""Generate the default PowerStacks executive presentation template.

Usage:
    python make_template.py [--out templates/powerstacks-executive.pptx]

Builds the eight prototype slides that build_deck.py fills in. Every
fillable shape has the name build_deck.py looks for (the template contract in
README.md). Style: near-black background with a faint network graphic, an
indigo title bar with the logo, slate cards with accent stripes, and status
colors (green, amber, red) applied by build_deck.py. All artwork is drawn
with shapes; the only image is the PowerStacks logo.
"""

import argparse
import math
import os
import random

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "templates", "powerstacks-logo.png")

BG = RGBColor(0x0A, 0x0B, 0x11)
CARD = RGBColor(0x18, 0x1D, 0x30)
CARD_LINE = RGBColor(0x35, 0x44, 0x68)
ART_LINE = RGBColor(0x21, 0x24, 0x2F)
ART_DOT = RGBColor(0x23, 0x25, 0x2C)
BAR = RGBColor(0x21, 0x27, 0x7D)
BAR_TOP = RGBColor(0x30, 0x25, 0x8F)
STRIPE = RGBColor(0x45, 0x7B, 0xF0)
LIGHT = RGBColor(0x6E, 0xC5, 0xF7)
WHITE = RGBColor(0xF8, 0xF9, 0xFF)
MUTED = RGBColor(0xB9, 0xC4, 0xD7)
FONT = "Aptos"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
TOP_STRIPE = Inches(0.12)
BAR_Y, BAR_H = TOP_STRIPE, Inches(0.78)
BODY_TOP = Inches(1.35)
FOOT_Y = Inches(7.02)


def rect(slide, name, left, top, width, height, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, left, top, width, height)
    s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def text(slide, name, left, top, width, height, value, size=16, color=WHITE, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, caps=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(0))
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value.upper() if caps else value
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return tb


def network_art(slide, seed):
    """Faint constellation on the right side, drawn with shapes (no images)."""
    rnd = random.Random(seed)
    pts = [(rnd.uniform(8.6, 13.1), rnd.uniform(1.2, 7.2)) for _ in range(26)]
    for i, (x1, y1) in enumerate(pts):
        near = sorted(((math.hypot(x2 - x1, y2 - y1), j) for j, (x2, y2) in enumerate(pts) if j != i))[:2]
        for d, j in near:
            if j > i and d < 1.6:
                x2, y2 = pts[j]
                ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
                ln.name = "art_line"
                ln.line.color.rgb = ART_LINE
                ln.line.width = Pt(0.75)
    for x, y in pts:
        r = rnd.uniform(0.04, 0.11)
        dot = rect(slide, "art_dot", Inches(x - r), Inches(y - r), Inches(2 * r), Inches(2 * r), ART_DOT, MSO_SHAPE.OVAL)


def base(prs, seed):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    network_art(s, seed)
    return s


def title_bar(slide):
    """Top stripe, bar, and left accent share one geometry; logo and title are centered in the bar."""
    rect(slide, "bar_top", 0, 0, W, TOP_STRIPE, BAR_TOP)
    rect(slide, "title_bar", 0, BAR_Y, W, BAR_H, BAR)
    rect(slide, "bar_accent", 0, BAR_Y, TOP_STRIPE, BAR_H, STRIPE)
    logo_h = Inches(0.6)
    logo = slide.shapes.add_picture(LOGO, MARGIN - Inches(0.2), BAR_Y + (BAR_H - logo_h) // 2, height=logo_h)
    logo.name = "logo"
    left = logo.left + logo.width + Inches(0.3)
    text(slide, "slide_title", left, BAR_Y, W - left - MARGIN, BAR_H, "[Slide title]", size=24, bold=True,
         anchor=MSO_ANCHOR.MIDDLE, caps=True)


def footer(slide):
    text(slide, "page_number", MARGIN, FOOT_Y, Inches(0.5), Inches(0.25), "00", size=10, color=LIGHT, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, "source_footer", MARGIN + Inches(0.5), FOOT_Y, Inches(8.6), Inches(0.25),
         "[Source]  |  [Data collected: UTC]  |  [Scope]", size=10, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, "brand_footer", W - MARGIN - Inches(3.2), FOOT_Y, Inches(3.2), Inches(0.25), "POWERSTACKS", size=10,
         color=MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def card(slide, name, left, top, width, height, stripe_name=None, stripe_color=STRIPE, label=None):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    c.name = name
    c.adjustments[0] = 0.05
    c.fill.solid()
    c.fill.fore_color.rgb = CARD
    c.line.color.rgb = CARD_LINE
    c.line.width = Pt(0.75)
    c.shadow.inherit = False
    if label is not None:
        tf = c.text_frame
        tf.text = label
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(12), MUTED
    if stripe_name:
        rect(slide, stripe_name, left, top, Inches(0.07), height, stripe_color)
    return c


def table(slide, rows, cols, left, top, width, height, headers, body, widths):
    gf = slide.shapes.add_table(rows, cols, left, top, width, height)
    gf.name = "table"
    t = gf.table
    tblPr = t._tbl.tblPr
    for attr in ("bandRow", "firstRow"):
        tblPr.set(attr, "0")
    style = tblPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId")
    if style is not None:
        tblPr.remove(style)
    for c, w in enumerate(widths):
        t.columns[c].width = w
    for r in range(rows):
        for c in range(cols):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BAR if r == 0 else CARD
            cell.margin_left = cell.margin_right = Inches(0.14)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = (headers[c].upper() if r == 0 else body[c])
            run.font.name = FONT
            run.font.size = Pt(12 if r == 0 else 14)
            run.font.bold = r == 0
            run.font.color.rgb = WHITE if r == 0 else WHITE
    return gf


def build(out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    cw = W - 2 * MARGIN

    # 1. Cover
    s = base(prs, 1)
    rect(s, "bar_top", 0, 0, W, TOP_STRIPE, BAR_TOP)
    logo = s.shapes.add_picture(LOGO, 0, Inches(1.0), height=Inches(1.9))
    logo.left = (W - logo.width) // 2
    logo.name = "logo"
    text(s, "product", MARGIN, Inches(3.15), cw, Inches(0.45), "BI for Intune", size=20, color=LIGHT, bold=True, align=PP_ALIGN.CENTER)
    text(s, "presentation_title", MARGIN, Inches(3.65), cw, Inches(0.9), "[Presentation title]", size=40, bold=True,
         align=PP_ALIGN.CENTER, caps=True)
    rule = rect(s, "accent_rule", (W - Inches(1.6)) // 2, Inches(4.7), Inches(1.6), Inches(0.06), STRIPE)
    text(s, "audience", MARGIN, Inches(5.0), cw, Inches(0.45), "[Audience]", size=18, align=PP_ALIGN.CENTER)
    text(s, "collection_time", MARGIN, Inches(5.55), cw, Inches(0.35), "[Data collected: date and time UTC]", size=13, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, "scope", MARGIN, Inches(5.9), cw, Inches(0.35), "[Scope]", size=13, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, "brand_footer", W - MARGIN - Inches(3.2), FOOT_Y, Inches(3.2), Inches(0.25), "POWERSTACKS", size=10,
         color=MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    # 2. Executive summary
    s = base(prs, 2)
    title_bar(s)
    gap = Inches(0.3)
    tile_w = (cw - 2 * gap) // 3
    top = BODY_TOP + Inches(0.35)
    for i in range(3):
        left = MARGIN + i * (tile_w + gap)
        card(s, f"metric_card_{i}", left, top, tile_w, Inches(1.85), stripe_name=f"metric_status_{i}")
        text(s, f"metric_label_{i}", left + Inches(0.35), top + Inches(0.28), tile_w - Inches(0.6), Inches(0.4),
             f"[Metric {i + 1}]", size=13, color=MUTED, caps=True, bold=True)
        text(s, f"metric_value_{i}", left + Inches(0.35), top + Inches(0.78), tile_w - Inches(0.6), Inches(0.85),
             f"[Value {i + 1}]", size=34, bold=True)
    card(s, "summary_card", MARGIN, Inches(4.0), cw, Inches(2.6), stripe_name="summary_stripe", stripe_color=LIGHT)
    text(s, "summary", MARGIN + Inches(0.35), Inches(4.25), cw - Inches(0.7), Inches(1.4), "[Headline: largest finding and next deadline]", size=20)
    text(s, "summary_caveat", MARGIN + Inches(0.35), Inches(5.75), cw - Inches(0.7), Inches(0.7), "[Most material caveat]", size=13, color=MUTED)
    footer(s)

    # 3. Findings table
    s = base(prs, 3)
    title_bar(s)
    table(s, 5, 3, MARGIN, BODY_TOP + Inches(0.35), cw, Inches(3.4),
          ["Finding", "Affected devices", "Applicable devices / share"],
          ["[Finding]", "[Count]", "[Population / %]"],
          [int(cw * 0.56), int(cw * 0.18), cw - int(cw * 0.56) - int(cw * 0.18)])
    text(s, "ranking_note", MARGIN, Inches(6.1), cw, Inches(0.7), "[Ranking method, filters, overlap note]", size=12, color=MUTED)
    footer(s)

    # 4. Finding detail
    s = base(prs, 4)
    title_bar(s)
    text(s, "main_metric", MARGIN, BODY_TOP + Inches(0.2), Inches(7.4), Inches(0.75), "[Count (share)]", size=34, bold=True, color=LIGHT)
    card(s, "donut_placeholder", MARGIN, Inches(2.45), Inches(2.95), Inches(3.05), label="[Affected vs not affected]")
    card(s, "chart_placeholder", MARGIN + Inches(3.15), Inches(2.45), Inches(4.3), Inches(3.05), label="[Check-in chart]")
    side_l = MARGIN + Inches(7.85)
    side_w = W - side_l - MARGIN
    card(s, "narrative_card", side_l, BODY_TOP + Inches(0.2), side_w, Inches(5.0), stripe_name="narrative_stripe")
    y = BODY_TOP + Inches(0.4)
    for name, label in (("finding_observation", "What we found"), ("finding_implication", "Why it matters"), ("finding_investigation", "Next step")):
        text(s, name + "_label", side_l + Inches(0.3), y, side_w - Inches(0.5), Inches(0.3), label, size=12, color=LIGHT, bold=True, caps=True)
        text(s, name, side_l + Inches(0.3), y + Inches(0.32), side_w - Inches(0.5), Inches(1.15), f"[{label}]", size=14)
        y += Inches(1.6)
    text(s, "finding_caveat", MARGIN, Inches(5.7), Inches(7.45), Inches(0.8), "[Check-in context, status age, confidence]", size=12, color=MUTED)
    footer(s)

    # 5. Breakdown
    s = base(prs, 5)
    title_bar(s)
    card(s, "chart_placeholder", MARGIN, BODY_TOP + Inches(0.3), Inches(8.0), Inches(4.0), label="[Breakdown chart]")
    side_l = MARGIN + Inches(8.35)
    side_w = W - side_l - MARGIN
    card(s, "breakdown_card", side_l, BODY_TOP + Inches(0.3), side_w, Inches(4.0), stripe_name="breakdown_stripe", stripe_color=LIGHT)
    text(s, "breakdown_summary", side_l + Inches(0.3), BODY_TOP + Inches(0.55), side_w - Inches(0.5), Inches(1.9), "[Key observation]", size=18)
    text(s, "breakdown_context", side_l + Inches(0.3), Inches(3.8), side_w - Inches(0.5), Inches(1.6), "[Denominator and overlap context]", size=13, color=MUTED)
    text(s, "breakdown_scope", MARGIN, Inches(5.85), Inches(8.0), Inches(0.8), "[Category definitions and scope]", size=12, color=MUTED)
    footer(s)

    # 6. Deadlines
    s = base(prs, 6)
    title_bar(s)
    table(s, 4, 3, MARGIN, BODY_TOP + Inches(0.35), cw, Inches(2.6),
          ["Date", "What it represents", "Affected devices"],
          ["[Date]", "[Source-reported deadline]", "[Count / scope]"],
          [int(cw * 0.18), int(cw * 0.52), cw - int(cw * 0.18) - int(cw * 0.52)])
    text(s, "deadline_window", MARGIN, Inches(6.1), cw, Inches(0.7), "[Display window and next deadline beyond it]", size=12, color=MUTED)
    footer(s)

    # 7. Data notes
    s = base(prs, 7)
    title_bar(s)
    y = BODY_TOP + Inches(0.35)
    for i, (t, b) in enumerate((("Population and scope", "[Definition and exclusions]"),
                                ("Freshness and coverage", "[Collection time and missing data]"),
                                ("Confidence and limitations", "[Check-in age, status age, caveats]"))):
        card(s, f"note_card_{i}", MARGIN, y, cw, Inches(1.3), stripe_name=f"note_stripe_{i}")
        text(s, f"note_title_{i}", MARGIN + Inches(0.35), y + Inches(0.22), Inches(3.5), Inches(0.9), t, size=14, color=LIGHT, bold=True, caps=True)
        text(s, f"note_body_{i}", MARGIN + Inches(4.1), y + Inches(0.22), cw - Inches(4.4), Inches(0.9), b, size=14)
        y += Inches(1.5)
    footer(s)

    # 8. Next investigation
    s = base(prs, 8)
    title_bar(s)
    card(s, "next_card", MARGIN, BODY_TOP + Inches(0.35), cw, Inches(2.3), stripe_name="next_stripe", stripe_color=LIGHT)
    text(s, "investigation", MARGIN + Inches(0.4), BODY_TOP + Inches(0.6), cw - Inches(0.8), Inches(1.8), "[Recommended investigation]", size=24, bold=True)
    text(s, "supporting_context", MARGIN, Inches(4.3), cw, Inches(1.1), "[Supporting findings and limitations]", size=16, color=MUTED)
    text(s, "followup", MARGIN, Inches(5.6), cw, Inches(0.8), "[Offer to narrow scope or change ranking]", size=14, color=LIGHT)
    footer(s)

    cp = prs.core_properties
    cp.title, cp.author, cp.last_modified_by, cp.comments, cp.subject, cp.keywords = (
        "PowerStacks executive briefing template", "PowerStacks", "", "", "", "")
    prs.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", default=os.path.join(HERE, "templates", "powerstacks-executive.pptx"))
    build(ap.parse_args().out)
