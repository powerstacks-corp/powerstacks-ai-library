"""Build an executive presentation from a briefing data file.

Usage:
    python build_deck.py --data briefing.json --out briefing.pptx [--template custom.pptx]

The data file follows executive-data.schema.json. Without --template, the
PowerStacks template next to this script is used. A customer template is
used in one of two ways:

- Named-shape mode: the template contains shapes with the names used by the
  PowerStacks template (for example `presentation_title`, `metric_value_0`,
  `chart_placeholder`). Its slides are used as prototypes and filled in.
- Layout mode: any other template. All of its existing slides are removed and
  new slides are built from its slide layouts (title and content
  placeholders), so the customer's theme, fonts, and colors apply.

Charts are native PowerPoint charts, so they stay editable. The script only
reads the data file and the template and writes the output file.
"""

import argparse
import copy
import io
import json
import math
import os
import sys
import zipfile

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt
except ImportError:
    sys.exit("This script needs the python-pptx package (pip install python-pptx).")

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_TEMPLATE = os.path.join(HERE, "templates", "powerstacks-executive.pptx")

# Roles are recognised by a shape name that only that prototype slide has.
ROLE_MARKERS = [
    ("cover", "presentation_title"),
    ("summary", "metric_value_0"),
    ("finding_detail", "finding_observation"),
    ("breakdown", "breakdown_summary"),
    ("deadlines", "deadline_window"),
    ("notes", "note_title_0"),
    ("next", "investigation"),
    ("findings_table", "ranking_note"),
]

DEFAULT_HEADERS = {
    "findings_table": ["Finding", "Affected devices", "Applicable devices / share"],
    "deadlines": ["Date", "What it represents", "Affected devices"],
}

# Colors for the PowerStacks template (dark background).
PS_SERIES = ["457BF0", "6EC5F7", "B9C4D7", "F8F9FF"]
PS_TEXT = "F8F9FF"
PS_MUTED = "B9C4D7"
PS_GRID = "21242F"
STATUS_COLORS = {"good": "61CE70", "warn": "F5B041", "bad": "FF5A4F"}
DEFAULT_THRESHOLDS = {"good": 95, "warn": 85}


def finding_status(f, thresholds):
    """good / warn / bad from the share of relevant devices not affected, or None."""
    count, pop = f.get("count"), f.get("population")
    if not isinstance(count, int) or not isinstance(pop, int) or pop <= 0:
        return None
    ok = 100.0 * (1 - count / pop)
    return "good" if ok >= thresholds["good"] else "warn" if ok >= thresholds["warn"] else "bad"


def color_text(shape, hex_color):
    if shape is None or not shape.has_text_frame:
        return
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor.from_string(hex_color)


def color_fill(shape, hex_color):
    if shape is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(hex_color)


# ---------------------------------------------------------------- helpers

def load_data(path):
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    for key in ("title", "collectedUtc", "scope", "findings"):
        if key not in data:
            sys.exit(f"Data file is missing required field '{key}'.")
    return data


def shape_by_name(slide, name):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    return None


def slide_role(slide):
    names = {s.name for s in slide.shapes}
    for role, marker in ROLE_MARKERS:
        if marker in names:
            return role
    return None


def _font_size(shape):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size:
                return r.font.size.pt
    return 18.0


CAPS_SHAPES = {"slide_title", "presentation_title"}


def set_text(shape, text, autosize=True):
    """Replace a shape's text, keeping the first run's formatting."""
    if shape is None or not shape.has_text_frame:
        return
    text = "" if text is None else str(text)
    if shape.name in CAPS_SHAPES and getattr(shape, "_ps_caps", True):
        text = text.upper()
    tf = shape.text_frame
    p0 = tf.paragraphs[0]._p
    r0 = p0.find(qn("a:r"))
    rpr = copy.deepcopy(r0.find(qn("a:rPr"))) if r0 is not None and r0.find(qn("a:rPr")) is not None else None
    for p in list(tf._txBody.findall(qn("a:p")))[1:]:
        tf._txBody.remove(p)
    for child in list(p0):
        if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p0.remove(child)
    template_p = copy.deepcopy(p0)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = p0 if i == 0 else copy.deepcopy(template_p)
        if i > 0:
            tf._txBody.append(p)
        r = p.makeelement(qn("a:r"), {})
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        t = r.makeelement(qn("a:t"), {})
        t.text = line
        r.append(t)
        end = p.find(qn("a:endParaRPr"))
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)
    tf.word_wrap = True
    if autosize:
        fit_font(shape, text)


def fit_font(shape, text, min_pt=9.0, start_pt=None):
    """Shrink the font when the text would overflow the shape (estimate)."""
    size = start_pt or _font_size(shape)
    width_pt = Emu(shape.width).pt * 0.95
    height_pt = Emu(shape.height).pt * 0.95
    def lines_needed(pt):
        per_line = max(1, int(width_pt / (pt * 0.52)))
        return sum(max(1, math.ceil(len(line) / per_line)) for line in text.split("\n"))
    pt = size
    while pt > min_pt and lines_needed(pt) * pt * 1.25 > height_pt:
        pt -= 0.5
    if pt < size or start_pt:
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(pt)


def duplicate_slide(prs, src):
    """Append a copy of a slide (shapes, background, and picture links)."""
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.is_external:
            continue
        if rel.reltype.endswith("/image"):
            rid_map[rid] = new.part.relate_to(rel._target, rel.reltype)
    src_bg = src._element.cSld.find(qn("p:bg"))
    if src_bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(src_bg))
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        el = copy.deepcopy(el)
        for node in el.iter():
            for attr in (qn("r:embed"), qn("r:link"), qn("r:id")):
                if attr in node.attrib and node.attrib[attr] in rid_map:
                    node.attrib[attr] = rid_map[node.attrib[attr]]
        new.shapes._spTree.append(el)
    return new


def delete_slide(prs, slide):
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        if prs.part.related_part(sld_id.rId) is slide.part:
            prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)
            return


def fill_table(table, headers, rows):
    """Fill a table: header row, then data rows (rows are added or removed)."""
    tbl = table._tbl
    trs = tbl.findall(qn("a:tr"))
    proto = trs[-1]
    for tr in trs[1:]:
        tbl.remove(tr)
    for _ in rows:
        tbl.append(copy.deepcopy(proto))
    for c, h in enumerate(headers[: len(table.columns)]):
        set_cell(table.cell(0, c), h)
    for r, row in enumerate(rows, start=1):
        for c in range(len(table.columns)):
            set_cell(table.cell(r, c), row[c] if c < len(row) else "")


def set_cell(cell, text):
    tf = cell.text_frame
    p0 = tf.paragraphs[0]._p
    r0 = p0.find(qn("a:r"))
    rpr = copy.deepcopy(r0.find(qn("a:rPr"))) if r0 is not None and r0.find(qn("a:rPr")) is not None else None
    for p in list(tf._txBody.findall(qn("a:p")))[1:]:
        tf._txBody.remove(p)
    for child in list(p0):
        if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p0.remove(child)
    r = p0.makeelement(qn("a:r"), {})
    if rpr is not None:
        r.append(rpr)
    t = r.makeelement(qn("a:t"), {})
    t.text = "" if text is None else str(text)
    r.append(t)
    end = p0.find(qn("a:endParaRPr"))
    if end is not None:
        end.addprevious(r)
    else:
        p0.append(r)


def resize_table_rows(graphic_frame, max_height):
    rows = graphic_frame.table.rows
    if not len(rows):
        return
    each = int(min(max_height / len(rows), Emu(Pt(40))))
    for row in rows:
        row.height = each


def add_chart(slide, box, chart_spec, branded):
    """Add a native, editable bar chart in the box (left, top, width, height)."""
    cats = [str(c) for c in chart_spec.get("categories", [])]
    series = chart_spec.get("series", [])
    if not cats or not series:
        return None
    cd = CategoryChartData()
    cd.categories = cats
    for s in series:
        cd.add_series(s.get("name", ""), [v if v is not None else 0 for v in s.get("values", [])])
    kind = XL_CHART_TYPE.BAR_STACKED if chart_spec.get("stacked") else XL_CHART_TYPE.BAR_CLUSTERED
    left, top, width, height = box
    chart = slide.shapes.add_chart(kind, left, top, width, height, cd).chart
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    cat_ax = chart.category_axis
    cat_ax.reverse_order = True  # first category at the top
    cat_ax.has_major_gridlines = False
    val_ax = chart.value_axis
    val_ax.has_major_gridlines = True
    top = max((v or 0) for s in series for v in s.get("values", [])) if series else 0
    if top <= 10:
        val_ax.major_unit = 1
        val_ax.minimum_scale = 0
    val_ax.tick_labels.number_format = '#,##0'
    val_ax.tick_labels.number_format_is_linked = False
    title = chart_spec.get("title")
    chart.has_title = bool(title)
    if title:
        chart.chart_title.text_frame.text = title
        for para in chart.chart_title.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(12)
                run.font.bold = False
                if branded:
                    run.font.color.rgb = RGBColor.from_string(PS_MUTED)
    plot = chart.plots[0]
    plot.vary_by_categories = False
    plot.gap_width = 60
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = '#,##0'
    labels.number_format_is_linked = False
    if not chart_spec.get("stacked"):
        labels.position = XL_LABEL_POSITION.OUTSIDE_END
    if branded:
        chart.font.size = Pt(11)
        chart.font.color.rgb = RGBColor.from_string(PS_TEXT)
        for i, s in enumerate(chart.series):
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = RGBColor.from_string(PS_SERIES[i % len(PS_SERIES)])
        val_ax.major_gridlines.format.line.color.rgb = RGBColor.from_string(PS_GRID)
        val_ax.format.line.fill.background()
        cat_ax.format.line.color.rgb = RGBColor.from_string(PS_GRID)
        val_ax.tick_labels.font.color.rgb = RGBColor.from_string(PS_MUTED)
        cat_ax.tick_labels.font.color.rgb = RGBColor.from_string(PS_TEXT)
        labels.font.color.rgb = RGBColor.from_string(PS_TEXT)
    else:
        chart.font.size = Pt(11)
    return chart


def _place(slide, ph):
    """Box for a chart: inside a card placeholder (kept as the frame), or the placeholder's own box."""
    if ph.shape_type == 1:  # auto shape used as a card: keep it, clear its label, inset the chart
        ph.text_frame.text = ""
        pad = Emu(Pt(10))
        return (ph.left + pad, ph.top + pad, ph.width - 2 * pad, ph.height - 2 * pad), False
    return (ph.left, ph.top, ph.width, ph.height), True


def replace_with_chart(slide, placeholder_name, chart_spec, branded):
    ph = shape_by_name(slide, placeholder_name)
    if ph is None:
        return
    box, remove = _place(slide, ph)
    if chart_spec and add_chart(slide, box, chart_spec, branded) is not None:
        if remove:
            ph._element.getparent().remove(ph._element)
    else:
        set_text(ph, (chart_spec or {}).get("emptyText", ""))


def add_donut(slide, placeholder_name, f, status):
    """Native doughnut: affected (status color) versus not affected (green)."""
    ph = shape_by_name(slide, placeholder_name)
    if ph is None:
        return
    count, pop = f.get("count"), f.get("population")
    if not isinstance(count, int) or not isinstance(pop, int) or pop <= 0:
        set_text(ph, "Applicable population not established")
        return
    box, remove = _place(slide, ph)
    cd = CategoryChartData()
    cd.categories = ["Affected", "Not affected"]
    cd.add_series("Devices", [count, max(0, pop - count)])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, *box, cd).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(11)
    chart.legend.font.color.rgb = RGBColor.from_string(PS_TEXT)
    chart.has_title = False
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format = '#,##0'
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(12)
    plot.data_labels.font.bold = True
    plot.data_labels.font.color.rgb = RGBColor.from_string("0A0B11")
    colors = [STATUS_COLORS.get(status or "bad", STATUS_COLORS["bad"]), STATUS_COLORS["good"]]
    if status == "good":
        colors[0] = STATUS_COLORS["warn"]
    for i, point in enumerate(chart.plots[0].series[0].points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor.from_string(colors[i])
        point.format.line.color.rgb = RGBColor.from_string("181D30")
    if remove:
        ph._element.getparent().remove(ph._element)


def fmt_int(n):
    return f"{n:,}" if isinstance(n, int) else ("" if n is None else str(n))


def finding_rows(data):
    rows = []
    for f in data["findings"]:
        share = f.get("share")
        pop = f.get("population")
        right = fmt_int(pop) if pop is not None else "Not established"
        if share:
            right += f" ({share})"
        rows.append([f.get("title", ""), fmt_int(f.get("count")), right])
    return rows


def footer_text(data):
    parts = [data.get("source", ""), "Data collected: " + data["collectedUtc"], data["scope"]]
    return "  |  ".join(p for p in parts if p)


# ---------------------------------------------------------- named-shape mode

def build_named(prs, data):
    thresholds = data.get("thresholds") or DEFAULT_THRESHOLDS
    protos = {}
    for slide in prs.slides:
        role = slide_role(slide)
        if role and role not in protos:
            protos[role] = slide
    originals = list(prs.slides)
    out = []

    def from_proto(role):
        if role not in protos:
            return None
        s = duplicate_slide(prs, protos[role])
        out.append(s)
        set_text(shape_by_name(s, "source_footer"), footer_text(data))
        brand = shape_by_name(s, "brand_footer")
        if brand is not None:
            set_text(brand, ("POWERSTACKS  •  " + data.get("product", "")).upper().strip(" •"), autosize=False)
        set_text(shape_by_name(s, "slide_title"), "", autosize=False)
        return s

    s = from_proto("cover")
    if s:
        set_text(shape_by_name(s, "product"), data.get("product", ""))
        set_text(shape_by_name(s, "presentation_title"), data["title"])
        set_text(shape_by_name(s, "audience"), data.get("audience", ""))
        set_text(shape_by_name(s, "collection_time"), "Data collected: " + data["collectedUtc"])
        set_text(shape_by_name(s, "scope"), data["scope"])

    summ = data.get("summary", {})
    s = from_proto("summary")
    if s:
        set_text(shape_by_name(s, "slide_title"), summ.get("slideTitle", "Executive summary"))
        metrics = summ.get("metrics", [])
        for i in range(3):
            m = metrics[i] if i < len(metrics) else {"label": "", "value": ""}
            set_text(shape_by_name(s, f"metric_label_{i}"), m.get("label", ""))
            set_text(shape_by_name(s, f"metric_value_{i}"), m.get("value", ""))
            if m.get("status") in STATUS_COLORS:
                color_fill(shape_by_name(s, f"metric_status_{i}"), STATUS_COLORS[m["status"]])
        set_text(shape_by_name(s, "summary"), summ.get("text", ""))
        set_text(shape_by_name(s, "summary_caveat"), summ.get("caveat", ""))

    s = from_proto("findings_table")
    if s:
        set_text(shape_by_name(s, "slide_title"), data.get("findingsTitle", "Findings by scale"))
        for shp in s.shapes:
            if shp.has_table:
                fill_table(shp.table, data.get("findingsHeaders", DEFAULT_HEADERS["findings_table"]), finding_rows(data))
                resize_table_rows(shp, Emu(Pt(270)))
                for r, f in enumerate(data["findings"], start=1):
                    st = finding_status(f, thresholds)
                    if st:
                        for p in shp.table.cell(r, 2).text_frame.paragraphs:
                            for run in p.runs:
                                run.font.color.rgb = RGBColor.from_string(STATUS_COLORS[st])
                break
        set_text(shape_by_name(s, "ranking_note"), data.get("rankingNote", ""))

    for f in data["findings"]:
        s = from_proto("finding_detail")
        if not s:
            break
        set_text(shape_by_name(s, "slide_title"), f.get("title", ""))
        main = fmt_int(f.get("count")) + " " + f.get("countLabel", "devices")
        if f.get("share"):
            main += f" ({f['share']})"
        set_text(shape_by_name(s, "main_metric"), main)
        status = finding_status(f, thresholds)
        if status:
            color_text(shape_by_name(s, "main_metric"), STATUS_COLORS[status])
            color_fill(shape_by_name(s, "narrative_stripe"), STATUS_COLORS[status])
        add_donut(s, "donut_placeholder", f, status)
        set_text(shape_by_name(s, "finding_observation"), f.get("observation", ""))
        set_text(shape_by_name(s, "finding_implication"), f.get("implication", ""))
        set_text(shape_by_name(s, "finding_investigation"), f.get("nextStep", ""))
        set_text(shape_by_name(s, "finding_caveat"), f.get("caveat", ""))
        replace_with_chart(s, "chart_placeholder", f.get("chart"), branded=True)
        bd = f.get("breakdown")
        if bd:
            b = from_proto("breakdown")
            if b:
                set_text(shape_by_name(b, "slide_title"), bd.get("title", ""))
                set_text(shape_by_name(b, "breakdown_summary"), bd.get("summary", ""))
                set_text(shape_by_name(b, "breakdown_context"), bd.get("context", ""))
                set_text(shape_by_name(b, "breakdown_scope"), bd.get("scope", ""))
                replace_with_chart(b, "chart_placeholder", bd.get("chart"), branded=True)

    if data.get("deadlines") is not None:
        s = from_proto("deadlines")
        if s:
            set_text(shape_by_name(s, "slide_title"), data.get("deadlinesTitle", "Upcoming deadlines"))
            rows = [[d.get("date", ""), d.get("what", ""), d.get("scope", "")] for d in data["deadlines"]]
            if not rows:
                rows = [["", data.get("deadlinesEmptyText", ""), ""]]
            for shp in s.shapes:
                if shp.has_table:
                    fill_table(shp.table, data.get("deadlinesHeaders", DEFAULT_HEADERS["deadlines"]), rows)
                    resize_table_rows(shp, Emu(Pt(270)))
                    break
            set_text(shape_by_name(s, "deadline_window"), data.get("deadlineWindow", ""))

    notes = data.get("notes", [])
    if notes:
        s = from_proto("notes")
        if s:
            set_text(shape_by_name(s, "slide_title"), data.get("notesTitle", "Data notes"))
            for i in range(3):
                n = notes[i] if i < len(notes) else {"title": "", "body": ""}
                set_text(shape_by_name(s, f"note_title_{i}"), n.get("title", ""))
                set_text(shape_by_name(s, f"note_body_{i}"), n.get("body", ""))

    nxt = data.get("next")
    if nxt:
        s = from_proto("next")
        if s:
            set_text(shape_by_name(s, "slide_title"), nxt.get("slideTitle", "Next investigation"))
            set_text(shape_by_name(s, "investigation"), nxt.get("investigation", ""))
            set_text(shape_by_name(s, "supporting_context"), nxt.get("context", ""))
            set_text(shape_by_name(s, "followup"), nxt.get("followup", ""))

    for slide in originals:
        delete_slide(prs, slide)
    for i, s in enumerate(out, start=1):
        set_text(shape_by_name(s, "page_number"), f"{i:02d}", autosize=False)


# ---------------------------------------------------------------- layout mode

CONTENT_TYPES = ("BODY", "OBJECT", "SUBTITLE")
CHROME_TYPES = ("DATE", "FOOTER", "SLIDE_NUMBER")


def ph_type(ph):
    return str(ph.placeholder_format.type).split(" ")[0].split(".")[-1]


def pick_layout(prs, want_title_slide):
    """Choose a title layout, or a content layout with one large content area."""
    area = prs.slide_width * prs.slide_height
    best, best_score = None, None
    for layout in prs.slide_layouts:
        phs = list(layout.placeholders)
        types = [ph_type(ph) for ph in phs]
        name = layout.name.lower()
        content = [ph for ph in phs if ph_type(ph) in CONTENT_TYPES]
        if want_title_slide:
            score = (3 if "CENTER_TITLE" in types else 0) + (1 if "SUBTITLE" in types else 0)
            score += 1 if "title" in name else 0
        else:
            if "TITLE" not in types or not content:
                continue
            largest = max(ph.width * ph.height for ph in content) / area
            score = largest * 10 - 2 * (len(content) - 1)
            score += 1 if "content" in name else 0
            score -= 3 if "vertical" in name or "picture" in name else 0
        if best_score is None or score > best_score:
            best, best_score = layout, score
    return best or prs.slide_layouts[0]


def title_and_body(slide):
    """Return the title placeholder and the largest content placeholder."""
    title, body = None, None
    for ph in slide.placeholders:
        t = ph_type(ph)
        if title is None and t in ("TITLE", "CENTER_TITLE"):
            title = ph
        elif t in CONTENT_TYPES and (body is None or ph.width * ph.height > body.width * body.height):
            body = ph
    return title, body


def drop_unused_placeholders(slide, keep):
    for ph in list(slide.placeholders):
        if ph in keep or ph_type(ph) in CHROME_TYPES:
            continue
        ph._element.getparent().remove(ph._element)


def layout_slide(prs, layout, title, body_text=None):
    s = prs.slides.add_slide(layout)
    t, b = title_and_body(s)
    if t is not None:
        t.text = title
    box = None
    keep = [t]
    if b is not None and b.width * b.height > 0.2 * prs.slide_width * prs.slide_height:
        box = (b.left, b.top, b.width, b.height)
        if body_text is not None:
            b.text_frame.text = body_text
            b.text_frame.word_wrap = True
            fit_font(b, body_text, min_pt=10, start_pt=20)
            keep.append(b)
    drop_unused_placeholders(s, keep)
    if box is None:
        box = (int(prs.slide_width * 0.06), int(prs.slide_height * 0.22),
               int(prs.slide_width * 0.88), int(prs.slide_height * 0.62))
        if body_text is not None:
            tb = add_textbox(s, box, body_text, size=20)
            fit_font(tb, body_text, min_pt=10, start_pt=20)
    return s, box


def add_textbox(slide, box, text, size=14):
    tb = slide.shapes.add_textbox(*box)
    tb.text_frame.word_wrap = True
    tb.text_frame.text = text
    for p in tb.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
    return tb


def add_footer(prs, slide, data, page):
    h = int(prs.slide_height * 0.05)
    box = (int(prs.slide_width * 0.04), int(prs.slide_height - h * 1.2), int(prs.slide_width * 0.85), h)
    add_textbox(slide, box, footer_text(data), size=9)
    box = (int(prs.slide_width * 0.9), int(prs.slide_height - h * 1.2), int(prs.slide_width * 0.06), h)
    add_textbox(slide, box, f"{page:02d}", size=9)


def add_table(slide, box, headers, rows):
    left, top, width, height = box
    n_rows = len(rows) + 1
    shape = slide.shapes.add_table(n_rows, len(headers), left, top, width, min(height, Emu(Pt(28)) * n_rows))
    for c, h in enumerate(headers):
        shape.table.cell(0, c).text = h
    for r, row in enumerate(rows, start=1):
        for c in range(len(headers)):
            shape.table.cell(r, c).text = str(row[c]) if c < len(row) else ""
    for r in range(n_rows):
        for c in range(len(headers)):
            for p in shape.table.cell(r, c).text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(12)
    return shape


def build_layout(prs, data):
    for slide in list(prs.slides):
        delete_slide(prs, slide)
    title_layout = pick_layout(prs, True)
    content_layout = pick_layout(prs, False)
    page = 0

    def next_page(slide):
        nonlocal page
        page += 1
        add_footer(prs, slide, data, page)

    s = prs.slides.add_slide(title_layout)
    t, b = title_and_body(s)
    if t is not None:
        t.text = data["title"]
    sub = "\n".join(x for x in [data.get("product", ""), data.get("audience", ""),
                                "Data collected: " + data["collectedUtc"], data["scope"]] if x)
    if b is not None:
        b.text_frame.text = sub
        fit_font(b, sub, min_pt=10, start_pt=20)
        drop_unused_placeholders(s, [t, b])
    else:
        add_textbox(s, (int(prs.slide_width * 0.08), int(prs.slide_height * 0.6),
                        int(prs.slide_width * 0.84), int(prs.slide_height * 0.25)), sub)
    page += 1

    summ = data.get("summary", {})
    metrics = "\n".join(f"{m.get('label', '')}: {m.get('value', '')}" for m in summ.get("metrics", []))
    body = "\n".join(x for x in [metrics, summ.get("text", ""), summ.get("caveat", "")] if x)
    s, _ = layout_slide(prs, content_layout, summ.get("slideTitle", "Executive summary"), body)
    next_page(s)

    s, box = layout_slide(prs, content_layout, data.get("findingsTitle", "Findings by scale"))
    l, t_, w, h = box
    add_table(s, (l, t_, w, int(h * 0.8)), data.get("findingsHeaders", DEFAULT_HEADERS["findings_table"]), finding_rows(data))
    if data.get("rankingNote"):
        add_textbox(s, (l, t_ + int(h * 0.85), w, int(h * 0.15)), data["rankingNote"], size=11)
    next_page(s)

    for f in data["findings"]:
        s, box = layout_slide(prs, content_layout, f.get("title", ""))
        l, t_, w, h = box
        main = fmt_int(f.get("count")) + " " + f.get("countLabel", "devices") + (f" ({f['share']})" if f.get("share") else "")
        add_textbox(s, (l, t_, int(w * 0.55), int(h * 0.15)), main, size=24)
        if f.get("chart"):
            add_chart(s, (l, t_ + int(h * 0.18), int(w * 0.55), int(h * 0.62)), f["chart"], branded=False)
        side = "\n\n".join(x for x in [f.get("observation", ""), f.get("implication", ""), f.get("nextStep", "")] if x)
        add_textbox(s, (l + int(w * 0.6), t_, int(w * 0.4), int(h * 0.8)), side, size=13)
        if f.get("caveat"):
            add_textbox(s, (l, t_ + int(h * 0.84), w, int(h * 0.16)), f["caveat"], size=11)
        next_page(s)
        bd = f.get("breakdown")
        if bd:
            s, box = layout_slide(prs, content_layout, bd.get("title", ""))
            l, t_, w, h = box
            if bd.get("chart"):
                add_chart(s, (l, t_, int(w * 0.6), int(h * 0.8)), bd["chart"], branded=False)
            side = "\n\n".join(x for x in [bd.get("summary", ""), bd.get("context", "")] if x)
            add_textbox(s, (l + int(w * 0.64), t_, int(w * 0.36), int(h * 0.8)), side, size=13)
            if bd.get("scope"):
                add_textbox(s, (l, t_ + int(h * 0.84), w, int(h * 0.16)), bd["scope"], size=11)
            next_page(s)

    if data.get("deadlines") is not None:
        s, box = layout_slide(prs, content_layout, data.get("deadlinesTitle", "Upcoming deadlines"))
        l, t_, w, h = box
        rows = [[d.get("date", ""), d.get("what", ""), d.get("scope", "")] for d in data["deadlines"]] or [["", data.get("deadlinesEmptyText", ""), ""]]
        add_table(s, (l, t_, w, int(h * 0.8)), data.get("deadlinesHeaders", DEFAULT_HEADERS["deadlines"]), rows)
        if data.get("deadlineWindow"):
            add_textbox(s, (l, t_ + int(h * 0.85), w, int(h * 0.15)), data["deadlineWindow"], size=11)
        next_page(s)

    notes = data.get("notes", [])
    if notes:
        body = "\n\n".join(f"{n.get('title', '')}: {n.get('body', '')}" for n in notes)
        s, _ = layout_slide(prs, content_layout, data.get("notesTitle", "Data notes"), body)
        next_page(s)

    nxt = data.get("next")
    if nxt:
        body = "\n\n".join(x for x in [nxt.get("investigation", ""), nxt.get("context", ""), nxt.get("followup", "")] if x)
        s, _ = layout_slide(prs, content_layout, nxt.get("slideTitle", "Next investigation"), body)
        next_page(s)


# ---------------------------------------------------------------------- main

def open_template(path):
    """Open a .pptx, or a .potx/.potm by changing its main content type in memory."""
    if not path.lower().endswith((".potx", ".potm")):
        return Presentation(path)
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            blob = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                blob = blob.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
                blob = blob.replace(b"presentationml.template.macroEnabled.main+xml", b"presentationml.presentation.main+xml")
            dst.writestr(item, blob)
    buf.seek(0)
    return Presentation(buf)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--data", required=True, help="Briefing data file (JSON)")
    ap.add_argument("--out", required=True, help="Output .pptx path")
    ap.add_argument("--template", help="Customer template (.pptx or .potx). Default: PowerStacks template")
    args = ap.parse_args(argv)

    data = load_data(args.data)
    template = args.template or DEFAULT_TEMPLATE
    prs = open_template(template)
    named = any(slide_role(s) for s in prs.slides)
    if named:
        build_named(prs, data)
    else:
        build_layout(prs, data)
    cp = prs.core_properties
    cp.title = data["title"]
    cp.author = data.get("author", "")
    cp.last_modified_by = ""
    cp.comments = ""
    cp.subject = data.get("product", "")
    cp.keywords = ""
    prs.save(args.out)
    print(f"Wrote {args.out} ({len(prs.slides)} slides, {'named-shape' if named else 'layout'} mode)")


if __name__ == "__main__":
    main()
